"""Helpers for building the Krish dossier with python-docx (styles, tables, callouts, TOC, figures)."""
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor, Emu

FONT = "Arial"
MONO = "Consolas"
BODY = "374151"
C = dict(title="1E293B", h1="0F172A", h2="1E3A8A", h3="334155", navy="1E293B", zebra="F8FAFC",
         border="CBD5E1", code_bg="F1F5F9", muted="64748B")
CALLOUT = {
    "warn": ("FEF2F2", "DC2626", "Key finding"),
    "info": ("EFF6FF", "2563EB", "Architectural discovery"),
    "note": ("FFFBEB", "D97706", "Implementation note"),
    "context": ("F1F5F9", "1E3A8A", "The big picture"),
}
BADGE = {"green": "15803D", "red": "B91C1C", "amber": "B45309", "blue": "1D4ED8", "grey": "475569"}
CONTENT_CM = 17.0


def rgb(h):
    return RGBColor.from_string(h)


def _set_fonts(rpr_owner, name):
    rpr = rpr_owner.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), name)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]


# Schema order: children that must come AFTER the element being inserted.
_AFTER = {
    ("pPr", "pBdr"): ("w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
                      "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid",
                      "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
                      "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId",
                      "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange"),
    ("tcPr", "shd"): ("w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark",
                      "w:headers", "w:cellIns", "w:cellDel", "w:cellMerge", "w:tcPrChange"),
    ("tblPr", "tblBorders"): ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription", "w:tblPrChange"),
    ("tblPr", "tblCellMar"): ("w:tblLook", "w:tblCaption", "w:tblDescription", "w:tblPrChange"),
}
_AFTER[("pPr", "shd")] = _AFTER[("pPr", "pBdr")][1:]


def put(parent, el):
    """Insert el into a *Pr element at its schema position."""
    key = (parent.tag.split("}")[1], el.tag.split("}")[1])
    parent.insert_element_before(el, *_AFTER[key])
    return el


def _shd(parent_pr, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    put(parent_pr, shd)


class Dossier:
    def __init__(self):
        self.doc = Document()
        self.toc = []          # (level, text, bookmark)
        self.bm = 0
        self.fig = 0
        self.tab = 0
        self._setup()

    # ------------------------------------------------------------------ setup
    def _setup(self):
        d = self.doc
        s = d.sections[0]
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.orientation = WD_ORIENT.PORTRAIT
        s.left_margin = s.right_margin = Cm(2.0)
        s.top_margin, s.bottom_margin = Cm(2.0), Cm(1.8)
        st = d.styles
        n = st["Normal"]
        n.font.name, n.font.size, n.font.color.rgb = FONT, Pt(10.5), rgb(BODY)
        _set_fonts(n.element, FONT)
        n.paragraph_format.line_spacing = 1.25
        n.paragraph_format.space_after = Pt(4)
        n.paragraph_format.space_before = Pt(0)

        def head(name, size, color, bold=True, before=0, after=4, pbb=False):
            h = st[name]
            h.font.name, h.font.size, h.font.bold, h.font.italic = FONT, Pt(size), bold, False
            h.font.color.rgb = rgb(color)
            _set_fonts(h.element, FONT)
            pf = h.paragraph_format
            pf.space_before, pf.space_after = Pt(before), Pt(after)
            pf.keep_with_next = True
            pf.page_break_before = pbb
            pf.line_spacing = 1.1
            return h

        head("Title", 26, C["title"], after=6)
        t = st["Title"].element.get_or_add_pPr()
        for b in t.findall(qn("w:pBdr")):
            t.remove(b)
        head("Subtitle", 13, C["h2"], bold=False, after=10)
        head("Heading 1", 18, C["h1"], before=0, after=8, pbb=True)
        head("Heading 2", 14, C["h2"], before=14, after=5)
        head("Heading 3", 12, C["h3"], before=10, after=3)
        for name, ind in (("TOC 1", 0), ("TOC 2", 0.6)):
            try:
                ts = st[name]
            except KeyError:
                ts = st.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            ts.base_style = st["Normal"]
            ts.font.size = Pt(10.5 if ind == 0 else 10)
            ts.font.bold = ind == 0
            ts.paragraph_format.left_indent = Cm(ind)
            ts.paragraph_format.space_after = Pt(2 if ind else 3)
            ts.paragraph_format.space_before = Pt(4 if ind == 0 else 0)
        # header / footer
        hp = s.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = hp.add_run("Krish 9-Class YOLO26-seg  |  Technical Dossier v2.5")
        r.font.size, r.font.color.rgb = Pt(8), rgb(C["muted"])
        fp = s.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = fp.add_run("Page ")
        r.font.size, r.font.color.rgb = Pt(8), rgb(C["muted"])
        self._field(fp, "PAGE")
        r = fp.add_run(" of ")
        r.font.size, r.font.color.rgb = Pt(8), rgb(C["muted"])
        self._field(fp, "NUMPAGES")
        # ask Word to refresh fields (TOC page numbers) on open
        settings = d.settings.element
        uf = OxmlElement("w:updateFields")
        uf.set(qn("w:val"), "true")
        later = {"hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids", "mathPr", "attachedSchema",
                 "themeFontLang", "clrSchemeMapping", "doNotIncludeSubdocsInStats", "doNotAutoCompressPictures",
                 "forceUpgrade", "captions", "readModeInkLockDown", "smartTagType", "schemaLibrary", "shapeDefaults",
                 "doNotEmbedSmartTags", "decimalSymbol", "listSeparator"}
        nxt = next((c for c in settings if c.tag.split("}")[1] in later), None)
        if nxt is not None:
            nxt.addprevious(uf)
        else:
            settings.append(uf)
        zoom = settings.find(qn("w:zoom"))
        if zoom is not None and zoom.get(qn("w:percent")) is None:
            zoom.set(qn("w:percent"), "100")

    def _field(self, p, code):
        run = p.add_run()
        run.font.size, run.font.color.rgb = Pt(8), rgb(C["muted"])
        for kind, text in (("begin", None), ("instr", code), ("separate", None), ("text", "1"), ("end", None)):
            if kind == "instr":
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = f" {text} "
            elif kind == "text":
                el = OxmlElement("w:t")
                el.text = text
            else:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), kind)
            run._r.append(el)

    # ------------------------------------------------------------- inline text
    def runs(self, p, text, size=None, color=None, bold=False, italic=False):
        parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
        for part in parts:
            if not part:
                continue
            if part.startswith("**") and part.endswith("**"):
                r = p.add_run(part[2:-2])
                r.bold = True
            elif part.startswith("`") and part.endswith("`"):
                r = p.add_run(part[1:-1])
                r.font.name = MONO
                _set_fonts(r._r, MONO)
                r.font.color.rgb = rgb("0F172A")
                if size:
                    r.font.size = Pt(size - 0.5)
                else:
                    r.font.size = Pt(9.5)
                if bold:
                    r.bold = True
                if italic:
                    r.italic = True
                continue
            else:
                r = p.add_run(part)
                r.bold = bold or None
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = rgb(color)
            if italic:
                r.italic = True
        return p

    def p(self, text, align=None, size=None, color=None, italic=False, after=None, keep=False):
        para = self.doc.add_paragraph()
        self.runs(para, text, size=size, color=color, italic=italic)
        if align == "center":
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if after is not None:
            para.paragraph_format.space_after = Pt(after)
        if keep:
            para.paragraph_format.keep_with_next = True
        return para

    def bullets(self, items, level=1):
        style = "List Bullet" if level == 1 else "List Bullet 2"
        for it in items:
            if isinstance(it, (list, tuple)):
                self.bullets(it, level + 1)
                continue
            para = self.doc.add_paragraph(style=style)
            para.paragraph_format.space_after = Pt(2)
            self.runs(para, it)

    def numbered(self, items):
        for i, it in enumerate(items, 1):
            para = self.doc.add_paragraph()
            pf = para.paragraph_format
            pf.left_indent, pf.first_line_indent, pf.space_after = Cm(0.75), Cm(-0.75), Pt(3)
            pf.tab_stops.add_tab_stop(Cm(0.75))
            r = para.add_run(f"{i}.\t")
            r.bold = True
            self.runs(para, it)

    # ---------------------------------------------------------------- headings
    def h(self, level, text, new_page=True):
        text = text.replace("`", "")
        para = self.doc.add_heading("", level=level)
        if not new_page:
            para.paragraph_format.page_break_before = False
        name = None
        if level <= 2:
            self.bm += 1
            name = f"_Toc{self.bm:05d}"
            bs = OxmlElement("w:bookmarkStart")
            bs.set(qn("w:id"), str(self.bm))
            bs.set(qn("w:name"), name)
            para._p.append(bs)
        para.add_run(text)
        if name:
            be = OxmlElement("w:bookmarkEnd")
            be.set(qn("w:id"), str(self.bm))
            para._p.append(be)
            self.toc.append((level, text, name))
        return para

    def title_block(self, title, subtitle, meta_lines):
        t = self.doc.add_paragraph(style="Title")
        t.add_run(title)
        st = self.doc.add_paragraph(style="Subtitle")
        st.add_run(subtitle)
        rule = self.doc.add_paragraph()
        pPr = rule._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        b = OxmlElement("w:bottom")
        for k, v in (("w:val", "single"), ("w:sz", "12"), ("w:space", "1"), ("w:color", C["h2"])):
            b.set(qn(k), v)
        bdr.append(b)
        put(pPr, bdr)
        rule.paragraph_format.space_after = Pt(8)
        for line in meta_lines:
            self.p(line, size=10, color=C["h3"], after=2)

    def toc_placeholder(self):
        self.p("**Contents**", size=14, color=C["h2"], after=6)
        self._toc_anchor = self.doc.add_paragraph()
        return self._toc_anchor

    def build_toc(self):
        anchor = self._toc_anchor._p
        entries = []
        for i, (lvl, text, name) in enumerate(self.toc):
            para = self.doc.add_paragraph(style=f"TOC {lvl}")
            p = para._p
            if i == 0:
                r = OxmlElement("w:r")
                for kind in ("begin", "instr", "separate"):
                    if kind == "instr":
                        el = OxmlElement("w:instrText")
                        el.set(qn("xml:space"), "preserve")
                        el.text = ' TOC \\o "1-2" \\h \\z \\u '
                    else:
                        el = OxmlElement("w:fldChar")
                        el.set(qn("w:fldCharType"), kind)
                    r.append(el)
                p.append(r)
            hl = OxmlElement("w:hyperlink")
            hl.set(qn("w:anchor"), name)
            hl.set(qn("w:history"), "1")
            r = OxmlElement("w:r")
            rpr = OxmlElement("w:rPr")
            col = OxmlElement("w:color")
            col.set(qn("w:val"), C["h2"] if lvl == 1 else BODY)
            rpr.append(col)
            r.append(rpr)
            t = OxmlElement("w:t")
            t.set(qn("xml:space"), "preserve")
            t.text = text
            r.append(t)
            hl.append(r)
            p.append(hl)
            entries.append(p)
        end = OxmlElement("w:r")
        fc = OxmlElement("w:fldChar")
        fc.set(qn("w:fldCharType"), "end")
        end.append(fc)
        entries[-1].append(end)
        for p in entries:
            anchor.addprevious(p)
        anchor.getparent().remove(anchor)

    # ------------------------------------------------------------------ blocks
    def callout(self, kind, title, body):
        fill, line, label = CALLOUT[kind]
        head = f"**{label}: {title}**" if title else f"**{label}**"
        items = [head] + (body if isinstance(body, list) else [body])
        for i, text in enumerate(items):
            para = self.doc.add_paragraph()
            pf = para.paragraph_format
            pf.left_indent, pf.right_indent = Cm(0.25), Cm(0.1)
            pf.space_after = Pt(0) if i < len(items) - 1 else Pt(8)
            pf.space_before = Pt(6) if i == 0 else Pt(0)
            pf.keep_with_next = i < len(items) - 1
            pPr = para._p.get_or_add_pPr()
            bdr = OxmlElement("w:pBdr")
            for side, sz, colr, space in (("top", "4", fill, "4"), ("left", "32", line, "8"),
                                          ("bottom", "4", fill, "4"), ("right", "4", fill, "4")):
                b = OxmlElement(f"w:{side}")
                for k, v in (("w:val", "single"), ("w:sz", sz), ("w:space", space), ("w:color", colr)):
                    b.set(qn(k), v)
                bdr.append(b)
            put(pPr, bdr)
            _shd(pPr, fill)
            if text.startswith("- "):
                self.runs(para, "â€¢  " + text[2:], size=10)
                pf.left_indent = Cm(0.7)
                pf.first_line_indent = Cm(-0.35)
            else:
                self.runs(para, text, size=10, color=line if i == 0 else None)

    def code(self, lines, caption=None):
        if caption:
            self.p(f"**{caption}**", size=9, color=C["muted"], after=1, keep=True)
        if isinstance(lines, str):
            lines = lines.strip("\n").split("\n")
        for i, line in enumerate(lines):
            para = self.doc.add_paragraph()
            pf = para.paragraph_format
            pf.space_after = Pt(0) if i < len(lines) - 1 else Pt(8)
            pf.line_spacing = 1.0
            pf.left_indent = Cm(0.2)
            pf.keep_with_next = i < len(lines) - 1 and len(lines) < 25
            pPr = para._p.get_or_add_pPr()
            bdr = OxmlElement("w:pBdr")
            b = OxmlElement("w:left")
            for k, v in (("w:val", "single"), ("w:sz", "12"), ("w:space", "6"), ("w:color", "94A3B8")):
                b.set(qn(k), v)
            bdr.append(b)
            put(pPr, bdr)
            _shd(pPr, C["code_bg"])
            r = para.add_run(line if line else " ")
            r.font.name = MONO
            _set_fonts(r._r, MONO)
            r.font.size = Pt(8.5)
            r.font.color.rgb = rgb("1F2937")

    def table(self, headers, rows, align=None, widths=None, caption=None, size=9):
        n = len(headers)
        align = align or "l" * n
        widths = widths or [1] * n
        tot = sum(widths)
        wcm = [CONTENT_CM * w / tot for w in widths]
        if caption:
            self.tab += 1
            self.p(f"**Table {self.tab}.** {caption}", size=9, color=C["h3"], after=2, keep=True)
        t = self.doc.add_table(rows=1 + len(rows), cols=n)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        tblPr = t._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            b = OxmlElement(f"w:{side}")
            for k, v in (("w:val", "single"), ("w:sz", "4"), ("w:space", "0"), ("w:color", C["border"])):
                b.set(qn(k), v)
            borders.append(b)
        put(tblPr, borders)
        mar = OxmlElement("w:tblCellMar")
        for side, v in (("top", 50), ("left", 90), ("bottom", 50), ("right", 90)):
            e = OxmlElement(f"w:{side}")
            e.set(qn("w:w"), str(v))
            e.set(qn("w:type"), "dxa")
            mar.append(e)
        put(tblPr, mar)
        AL = {"l": WD_ALIGN_PARAGRAPH.LEFT, "r": WD_ALIGN_PARAGRAPH.RIGHT, "c": WD_ALIGN_PARAGRAPH.CENTER}
        for ri, row in enumerate([headers] + rows):
            tr = t.rows[ri]
            if ri == 0:
                trPr = tr._tr.get_or_add_trPr()
                th = OxmlElement("w:tblHeader")
                th.set(qn("w:val"), "true")
                trPr.append(th)
            cant = OxmlElement("w:cantSplit")
            tr._tr.get_or_add_trPr().append(cant)
            for ci in range(n):
                cell = tr.cells[ci]
                cell.width = Cm(wcm[ci])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcPr = cell._tc.get_or_add_tcPr()
                fill = C["navy"] if ri == 0 else (C["zebra"] if ri % 2 == 0 else "FFFFFF")
                _shd(tcPr, fill)
                val = row[ci] if ci < len(row) else ""
                para = cell.paragraphs[0]
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.05
                para.alignment = AL["l"] if ri == 0 and align[ci] == "l" else AL[align[ci]]
                if ri == 0:
                    r = para.add_run(str(val))
                    r.bold, r.font.size, r.font.color.rgb = True, Pt(size), rgb("FFFFFF")
                elif isinstance(val, tuple) and val[0] == "badge":
                    r = para.add_run(val[1])
                    r.bold, r.font.size, r.font.color.rgb = True, Pt(size - 0.5), rgb(BADGE[val[2]])
                else:
                    lines = str(val).split("\n")
                    for li, line in enumerate(lines):
                        if li:
                            para = cell.add_paragraph()
                            para.paragraph_format.space_after = Pt(0)
                            para.paragraph_format.line_spacing = 1.05
                            para.alignment = AL[align[ci]]
                        self.runs(para, line, size=size)
        # grid widths for Google Docs
        grid = t._tbl.tblGrid
        for gc, w in zip(grid.findall(qn("w:gridCol")), wcm):
            gc.set(qn("w:w"), str(int(w * 567)))
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def figure(self, path, caption, width_cm=15.5, root=None):
        self.fig += 1
        full = Path(root) / path if root else Path(path)
        if full.exists():
            self.doc.add_picture(str(full), width=Cm(width_cm))
            pic = self.doc.paragraphs[-1]
            pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pic.paragraph_format.keep_with_next = True
            pic.paragraph_format.space_after = Pt(2)
        else:
            para = self.doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pPr = para._p.get_or_add_pPr()
            _shd(pPr, "F1F5F9")
            self.runs(para, f"[Figure placeholder: insert `{path}`]", size=9, color=C["muted"])
            para.paragraph_format.keep_with_next = True
        cap = self.doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.runs(cap, f"**Figure {self.fig}.** {caption}", size=9, color=C["h3"])
        self.runs(cap, f"  Source: `{path}`", size=8.5, color=C["muted"])
        cap.paragraph_format.space_after = Pt(10)

    def big_picture(self, *paras):
        """Narrative lead-in that opens every chapter."""
        self.callout("context", "", list(paras))

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def save(self, path):
        if hasattr(self, "_toc_anchor"):
            self.build_toc()
        self.doc.save(path)
